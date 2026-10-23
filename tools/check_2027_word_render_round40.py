import bisect,json,re,sys
from pathlib import Path
import pypdfium2 as pdf
from docx import Document
import argparse
p=argparse.ArgumentParser();p.add_argument('render',type=Path);a=p.parse_args()
repo=Path(__file__).resolve().parents[1];root=repo/'publishing/2027/round40';render=a.render
d=pdf.PdfDocument(render/'2027_review_manuscript.pdf');pages=[];limits=[];total=0
norm=lambda s:re.sub(r'\s+','',s)
for i in range(len(d)):
 p=d[i];t=p.get_textpage();s=t.get_text_range();w,h=p.get_size();outside=[]
 for k in range(t.count_chars()):
  c=t.get_text_range(k,1)
  if not c.strip():continue
  l,b,r,top=t.get_charbox(k)
  if l<-.5 or b<-.5 or r>w+.5 or top>h+.5:outside.append(dict(character=c,box=[l,b,r,top]))
 text=norm(s);pages.append(dict(page=i+1,characters=len(text),japanese=len(re.findall(r'[\u3040-\u30ff\u4e00-\u9fff]',text)),outside=outside,has_replacement='\ufffd' in s,text=text))
 total+=len(text);limits.append(total)
alltext=''.join(p['text'] for p in pages)
old=Document(repo/'publishing/2027/baseline/2027_layout_baseline.docx');new=Document(root/'2027_review_manuscript.docx');diff=[dict(index=i,after=b.text) for i,(a,b) in enumerate(zip(old.paragraphs,new.paragraphs)) if a.text!=b.text];affected=[]
for e in diff:
 needle=norm(e['after']);pos=alltext.find(needle)
 if pos<0:affected.append(dict(paragraph=e['index'],text_found=False,text=e['after']));continue
 affected.append(dict(paragraph=e['index'],text_found=True,pages=list(range(bisect.bisect_right(limits,pos)+1,bisect.bisect_left(limits,pos+len(needle))+2))))
word=Document(root/'2027_review_manuscript.docx');major=[]
for p in word.paragraphs:
 if p.style is not None and p.style.name=='Heading 1':
  needle=norm(p.text);hits=[x for x in pages if needle in x['text']];starts=[x['page'] for x in hits if x['text'].startswith(needle)]
  major.append(dict(heading=p.text,page_starts=starts,word_page_break_before=p.paragraph_format.page_break_before))
report=dict(pages=len(d),render_font_environment='Explicit Normal style and Hiragino fonts;10.5pt body,14.5pt fixed line, heading/title automatic line height;400-page final render distinct from historical373',empty_pages=[x['page'] for x in pages if not x['characters']],pages_without_japanese=[x['page'] for x in pages if not x['japanese']],out_of_page_characters=sum(len(x['outside']) for x in pages),replacement_pages=[x['page'] for x in pages if x['has_replacement']],affected_paragraphs=affected,affected_pages=sorted({pg for x in affected if x['text_found'] for pg in x['pages']}),major_sections=major,full_page_visual_qa_completed=False,visual_method="This script checks geometry/text only; manually inspect all pages before setting visual approval")
(root/'render_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k!='affected_paragraphs'},ensure_ascii=False,indent=2))
