"""Run from task workspace parent of work/kindle; local rendering paths in this checkpoint."""
import bisect,json,re,sys
from pathlib import Path
import pypdfium2 as pdf
from docx import Document
root=Path('work/kindle/publishing/2027/round38');render=Path('work/render38_fonts')
d=pdf.PdfDocument(render/'2027_round38_review.pdf');pages=[];limits=[];total=0
norm=lambda s:re.sub(r'\s+','',s)
for i in range(len(d)):
 p=d[i];t=p.get_textpage();s=t.get_text_range();w,h=p.get_size();outside=[]
 for k in range(t.count_chars()):
  c=t.get_text_range(k,1)
  if not c.strip():continue
  l,b,r,top=t.get_charbox(k)
  if l<-.5 or b<-.5 or r>w+.5 or top>h+.5:outside.append(dict(character=c,box=[l,b,r,top]))
 text=norm(s);pages.append(dict(page=i+1,characters=len(text),japanese=len(re.findall(r'[\u3040-\u30ff\u4e00-\u9fff]',text)),outside=outside,has_replacement='\ufffd' in s,text=text))
 total+=len(text);limits.append(total)
alltext=''.join(p['text'] for p in pages)
diff=json.loads((root/'word_paragraph_diff.json').read_text());affected=[]
for e in diff:
 needle=norm(e['after']);pos=alltext.find(needle)
 if pos<0:affected.append(dict(paragraph=e['index'],text_found=False,text=e['after']));continue
 affected.append(dict(paragraph=e['index'],text_found=True,pages=list(range(bisect.bisect_right(limits,pos)+1,bisect.bisect_left(limits,pos+len(needle))+2))))
word=Document(root/'2027_round38_review.docx');major=[]
for p in word.paragraphs:
 if p.style is not None and p.style.name=='Heading 1':
  needle=norm(p.text);hits=[x for x in pages if needle in x['text']];starts=[x['page'] for x in hits if x['text'].startswith(needle)]
  major.append(dict(heading=p.text,page_starts=starts,word_page_break_before=p.paragraph_format.page_break_before))
report=dict(pages=len(d),render_font_environment='explicit fontconfig with macOS Hiragino fallback; original document has no Normal style/font defaults; pagination not comparable to historical 373-page rendering',empty_pages=[x['page'] for x in pages if not x['characters']],pages_without_japanese=[x['page'] for x in pages if not x['japanese']],out_of_page_characters=sum(len(x['outside']) for x in pages),replacement_pages=[x['page'] for x in pages if x['has_replacement']],affected_paragraphs=affected,affected_pages=sorted({pg for x in affected if x['text_found'] for pg in x['pages']}),major_sections=major,full_page_visual_qa_completed=False)
(root/'render_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k!='affected_paragraphs'},ensure_ascii=False,indent=2))
