#!/usr/bin/env python3
import hashlib,json,re,zipfile
from pathlib import Path
from lxml import etree as ET
from docx import Document
from docx.oxml.ns import qn
from audit_2027_epub import read_epub,norm,ID,txt
ROOT=Path(__file__).resolve().parents[1]
DIR=ROOT/'publishing/2027/round38'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def docx_records(path):
 d=Document(path);ps=d.paragraphs;marks={}
 for i,p in enumerate(ps):
  for bm in p._element.iter(qn('w:bookmarkStart')):
   id=bm.get(qn('w:name'),'')
   if ID.fullmatch(id):assert id not in marks;marks[id]=i
 qs=[];nodes={}
 for id,start in marks.items():
  if not id.startswith('Q_'):continue
  stem=[];sn=[];options=[];on=[]
  for p in ps[start+1:]:
   t=p.text;m=re.match(r'^([1-4])[.．]\s*(.*)',t,re.S)
   if m:
    assert int(m[1])==len(options)+1,id
    options.append(m[2]);on.append(p)
   elif not options:stem.append(t);sn.append(p)
   if len(options)==4:break
  ap=ps[marks['A_'+id[2:]]+1:marks['A_'+id[2:]]+6]
  m=re.fullmatch(r'正解[：:]\s*([1-4])',ap[0].text);assert m,id
  answer=int(m[1]);ex=[]
  for k,p in enumerate(ap[1:],1):
   m=re.fullmatch(r'([1-4])（(正|誤)）\s*(.*)',p.text,re.S)
   assert m and int(m[1])==k and (m[2]=='正')==(k==answer),id
   ex.append(m[3])
  qs.append(dict(id=id,stem='\n'.join(stem),options=options,answer=answer,explanations=ex));nodes[id]=dict(stem=sn,options=on,explanations=ap[1:])
 assert len(qs)==500 and len(marks)==1000
 return d,qs,nodes
def equivalent(a,b):
 return a['id']==b['id'] and a['answer']==b['answer'] and norm(a['stem'])==norm(b['stem']) and all([norm(x) for x in a[f]]==[norm(x) for x in b[f]] for f in ['options','explanations'])
def rewrite_word(p,text,label=None):
 assert not p._element.findall(qn('w:hyperlink'))
 for r in list(p._element):
  if r.tag==qn('w:r'):p._element.remove(r)
 if label:
  p.add_run(label).bold=True;p.add_run('　'+text)
 else:p.add_run(text)
def rewrite_xml(p,text,label=None):
 assert not p.xpath('.//*[local-name()="a"]')
 for c in list(p):p.remove(c)
 p.text=None
 if label:
  s=ET.SubElement(p,'{http://www.w3.org/1999/xhtml}strong');s.text=label;s.tail='　'+text
 else:p.text=text
def main():
 ledger=json.loads((DIR/'revisions_14.json').read_text()); ep=ROOT/'publishing/2027/baseline/2027_layout_baseline.epub';wp=ep.with_suffix('.docx')
 assert sha(ep)==ledger['baseline_epub_sha256'] and sha(wp)==ledger['baseline_docx_sha256']
 z,roots,eq,en,qa=read_epub(ep);d,wq,wn=docx_records(wp)
 E={q['id']:q for q in eq};W={q['id']:q for q in wq};assert set(E)==set(W) and all(equivalent(E[id],W[id]) for id in E),'baseline DOCX/EPUB mismatch'
 for e in ledger['edits']:
  id=e['id'];before=e['before'];after=e['after'];assert E[id]==before and equivalent(W[id],before)
  for field in ['stem','options','explanations']:
   if before[field]==after[field]:continue
   vals=[after[field]] if field=='stem' else after[field]
   assert len(wn[id][field])==len(en[id][field])==len(vals),(id,field)
   for k,(v,p,x) in enumerate(zip(vals,wn[id][field],en[id][field]),1):
    label=f'{k}（'+('正' if k==after['answer'] else '誤')+'）' if field=='explanations' else None
    if field=='options':v=f'{k}. '+v
    rewrite_word(p,v,label);rewrite_xml(x,v,label)
 # Hard page-start rule: inherit and explicitly ensure all original major headings.
 major=[]
 for i,p in enumerate(d.paragraphs):
  if p.style is not None and p.style.name=='Heading 1':p.paragraph_format.page_break_before=True;major.append(dict(paragraph=i,text=p.text))
 outw=DIR/'2027_round38_review.docx';oute=DIR/'2027_round38_review.epub';d.save(outw)
 modified={n:ET.tostring(r,encoding='utf-8',xml_declaration=True) for n,r in roots.items() if z.read(n)!=ET.tostring(r,encoding='utf-8',xml_declaration=True)}
 # Preserve unedited XML byte-for-byte; only replace actual changed paragraph documents.
 changed_names=set()
 for e in ledger['edits']:
  id=e['id']
  for field in ['stem','options','explanations']:
   if e['before'][field]!=e['after'][field]:
    for p in en[id][field]:
     for n,r in roots.items():
      if p.getroottree().getroot() is r:changed_names.add(n)
 with zipfile.ZipFile(oute,'w') as oz:
  for info in z.infolist():oz.writestr(info,modified[info.filename] if info.filename in changed_names else z.read(info.filename))
 _,_,nq,_,nqa=read_epub(oute);nd,nwq,_=docx_records(outw);N={q['id']:q for q in nq};NW={q['id']:q for q in nwq}
 expected={q['id']:q for q in eq}
 for e in ledger['edits']:expected[e['id']]=e['after']
 assert N==expected and all(equivalent(N[id],NW[id]) for id in N)
 assert all(E[id]['answer']==N[id]['answer'] for id in E)
 old=Document(wp);diff=[dict(index=i,before=a.text,after=b.text) for i,(a,b) in enumerate(zip(old.paragraphs,nd.paragraphs)) if a.text!=b.text]
 assert len(old.paragraphs)==len(nd.paragraphs) and len(old.tables)==len(nd.tables)
 oldxml=ET.fromstring(zipfile.ZipFile(wp).read('word/document.xml'));newxml=ET.fromstring(zipfile.ZipFile(outw).read('word/document.xml'))
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 for query in ['//w:bookmarkStart/@w:name','//w:hyperlink/@w:anchor']:
  assert oldxml.xpath(query,namespaces=ns)==newxml.xpath(query,namespaces=ns)
 for a,b in zip(old.tables,nd.tables):assert ET.tostring(a._element)==ET.tostring(b._element)
 links=newxml.xpath('//w:hyperlink/@w:anchor',namespaces=ns);bookmarks=set(newxml.xpath('//w:bookmarkStart/@w:name',namespaces=ns));assert set(links)<=bookmarks
 assert not nqa['broken_hrefs'] and not nqa['duplicate_ids'] and nqa['replacement_characters']==0
 report=dict(baseline_docx_epub_full_field_match=500,updated_docx_epub_full_field_match=500,changed_questions=len(ledger['edits']),question_total=500,answer_keys_changed=0,changed_word_paragraphs=len(diff),word_paragraphs=len(nd.paragraphs),tables=len(nd.tables),word_internal_links=len(links),word_missing_destinations=0,major_page_starts=major,epub_modified_documents=sorted(changed_names),epub_qa=nqa,docx_sha256=sha(outw),epub_sha256=sha(oute),full_visual_qa_completed=False,epubcheck_completed=False,kindle_previewer_completed=False,release_ready=False)
 (DIR/'regression.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');(DIR/'word_paragraph_diff.json').write_text(json.dumps(diff,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
