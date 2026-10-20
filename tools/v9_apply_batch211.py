#!/usr/bin/env python3
"""Apply the five source-verified pending glossary edits and record batch 211.

This is a cumulative, one-time recording step.  The actual DOCX is rebuilt from
the immutable v8 source by tools/v9_rebuild.py after the ordered changes are
saved.  Historical evidence files are not rewritten: they remain accurate
records that the proposals were still pending when those checkpoints were made.
"""
from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "audit/v9"
EVIDENCE = OUT / "evidence"
DOCX = OUT / "登録日本語教員試験決定版問題集_v9_監査作業中.docx"
EXPECTED_PREVIOUS_DOCX = "190e6d900fb6331155c35a515b5118b88766e3f594d64eaebda1d3982fdc648c"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_proposals() -> list[dict]:
    batch208 = json.loads((EVIDENCE / "batch208_source_review.json").read_text())
    batch209 = json.loads((EVIDENCE / "batch209_source_review.json").read_text())
    batch210 = json.loads((EVIDENCE / "batch210_source_review.json").read_text())
    proposals = batch208["proposals"] + [batch209["proposal"], batch210["proposal"]]
    assert [p["term"] for p in proposals] == [
        "役割語", "連濁", "バーンアウト", "リキャストとプロンプト", "ディクトグロス"
    ]
    return proposals


def append_report(path: Path, text: str) -> None:
    old = path.read_text()
    marker = "## 第211チェックポイント 2026-10-21"
    assert marker not in old
    path.write_text(old.rstrip() + "\n\n\n" + text.strip() + "\n")


def main() -> None:
    assert sha256(DOCX) == EXPECTED_PREVIOUS_DOCX
    proposals = load_proposals()
    before_doc = Document(DOCX)
    before_paragraphs = [p.text for p in before_doc.paragraphs]
    for p in proposals:
        assert before_paragraphs.count(p["before"]) == 1, (p["term"], "before")
        assert before_paragraphs.count(p["after"]) == 0, (p["term"], "after")

    changes_path = OUT / "changes.json"
    changes = json.loads(changes_path.read_text())
    batch_changes = []
    for p in proposals:
        assert not any(c.get("before") == p["before"] and c.get("after") == p["after"] for c in changes)
        c = {
            "id": "APPENDIX_GLOSSARY",
            "field": "glossary",
            "index": None,
            "before": p["before"],
            "after": p["after"],
            "reason": p["reason"],
            "sources": p["sources"],
            "applied": True,
            "before_version": "第210チェックポイント時点の第206累積作業稿",
            "glossary_head": p["glossary_head"],
        }
        changes.append(c)
        batch_changes.append(c)

    changes_path.write_text(json.dumps(changes, ensure_ascii=False, indent=2) + "\n")
    (OUT / "changes_batch211.json").write_text(
        json.dumps(batch_changes, ensure_ascii=False, indent=2) + "\n"
    )

    previous_copy = EVIDENCE / "batch211_previous_working.docx"
    shutil.copy2(DOCX, previous_copy)
    subprocess.run([sys.executable, str(ROOT / "tools/v9_rebuild.py")], check=True, cwd=ROOT / "tools")

    after_doc = Document(DOCX)
    after_paragraphs = [p.text for p in after_doc.paragraphs]
    assert len(before_paragraphs) == len(after_paragraphs)
    changed = [(i, a, b) for i, (a, b) in enumerate(zip(before_paragraphs, after_paragraphs)) if a != b]
    assert len(changed) == 5
    assert sorted(a for _, a, _ in changed) == sorted(p["before"] for p in proposals)
    assert sorted(b for _, _, b in changed) == sorted(p["after"] for p in proposals)

    # The five changes are appendix-only: every question and answer object must
    # remain byte-for-byte equivalent after parsing.
    sys.path.insert(0, str(ROOT / "tools"))
    import quiz_audit as qa

    old_issues, old_q, old_a = qa.audit(str(previous_copy))
    new_issues, new_q, new_a = qa.audit(str(DOCX))
    assert old_q == new_q
    assert old_a == new_a
    assert old_issues == new_issues

    postfix = json.loads((OUT / "postfix_machine.json").read_text())
    assert postfix["questions"] == postfix["answers"] == 500
    assert postfix["replacements"] == 4321
    assert postfix["changed_questions"] == 493
    assert postfix["v8_unchanged"] is True
    assert postfix["answers_unchanged"] is True
    assert postfix["full_visual_qa_completed"] is False
    assert postfix["release_ready"] is False

    evidence = {
        "batch": 211,
        "date": "2026-10-21",
        "base_remote_commit": "f8520fe64b918d5cf25911c1a42a31aec55ed908",
        "scope": "適用待ちだった出典照合済み用語集5項目を累積稿へ反映し、500問本文・正答・機械警告の不変性を検証。",
        "applied_terms": [p["term"] for p in proposals],
        "sources": [p["sources"] for p in proposals],
        "changes_added": 5,
        "cumulative_changes": postfix["replacements"],
        "modified_questions": postfix["changed_questions"],
        "previous_docx_sha256": EXPECTED_PREVIOUS_DOCX,
        "current_docx_sha256": sha256(DOCX),
        "paragraph_differences": [
            {"paragraph_index": i, "before": a, "after": b}
            for i, a, b in changed
        ],
        "question_objects_unchanged": old_q == new_q,
        "answer_objects_unchanged": old_a == new_a,
        "machine_issues_unchanged": old_issues == new_issues,
        "raw_machine_warnings": len(new_issues),
        "undisposed_machine_warnings": 0,
        "v8_sha256": sha256(ROOT / "登録日本語教員試験決定版問題集_v8_修正版.docx"),
        "v8_expected_sha256": "4aef79631da7096e5837dd31473978d40d4932ba3cf8473c52138662a843d498",
        "v8_unchanged": True,
        "render_pending": True,
        "changed_pages_visual_qa_completed": False,
        "full_visual_qa_completed": False,
        "release_ready": False,
        "remaining": [
            "新稿の完全日本語フォントによる描画と変更頁画像確認",
            "用語集・前付け・ケース・索引の残る横断点検",
            "改頁調整",
            "編集後の最終機械検査",
            "最終全頁画像QA",
        ],
    }
    assert evidence["v8_sha256"] == evidence["v8_expected_sha256"]
    (EVIDENCE / "batch211_apply_review.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2) + "\n"
    )

    # The retained previous DOCX was only needed for deterministic comparison.
    previous_copy.unlink()

    append_report(
        OUT / "README.md",
        """## 第211チェックポイント 2026-10-21

作業環境の復旧後、適用待ちだった用語集5項目（役割語・連濁・バーンアウト・リキャストとプロンプト・ディクトグロス）を順序付き変更としてchanges.jsonへ追加し、v8不変原稿から累積再構成した。累積4321箇所、493問。500問の設問・全肢・正答・全解説と生機械警告は編集前後で同一。新稿の描画・変更頁画像QA、横断点検、改頁調整、最終機械検査、最終全頁QAは未完了。evidence/batch211_apply_review.json参照。""",
    )

    report = """## 第211チェックポイント 2026-10-21

作業環境の復旧を確認し、過去チェックポイントで未適用だった用語集5項目を、現在の原文との完全一致を確認してからchanges.jsonへ順番に追加した。対象は役割語、連濁、バーンアウト、リキャストとプロンプト、ディクトグロス。いずれも既取得の根拠本文に基づく用語集の横断整合修正であり、500問の設問・選択肢・正答・解説は変更していない。

v8のSHA256が4aef79631da7096e5837dd31473978d40d4932ba3cf8473c52138662a843d498で不変であることを再確認し、tools/v9_rebuild.pyで累積稿を安全に再構成した。編集前後の500問オブジェクト・500解答オブジェクト・生機械警告リストを直接比較し、全て同一。累積変更は4321箇所、変更対象は493問のまま。生警告159件の既存個別処置は未処置0件を維持するが、描画後の再確認を要する。

この時点では新稿の描画・変更頁画像QA、用語集・前付け・ケース・索引の残る横断点検、改頁調整、編集後の最終機械検査、最終全頁画像QAは未完了。最終版未制作、出版不可。詳細はevidence/batch211_apply_review.json。"""
    append_report(OUT / "critical_findings.md", report)
    append_report(OUT / "interim_report.md", report)
    append_report(
        OUT / "sources.md",
        """## 第211チェックポイント 用語集5項目の適用

第207～210チェックポイントで取得・照合したS201、S301、S193、S126、S104の根拠に基づき、役割語・連濁・バーンアウト・リキャストとプロンプト・ディクトグロスの5項目を累積稿へ反映した。新しい資料承認ではなく、各チェックポイントに保存した本文範囲・限界を維持する。全5項目はchanges_batch211.jsonとevidence/batch211_apply_review.jsonに対応し、500問本文・正答・解説は不変。""",
    )

    manifest_path = OUT / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    assert manifest["questions"] == manifest["answers"] == 500
    manifest["full_visual_qa_completed"] = False
    manifest["release_ready"] = False
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")

    print(json.dumps({
        "batch": 211,
        "changes_added": 5,
        "cumulative_changes": postfix["replacements"],
        "docx_sha256": sha256(DOCX),
        "raw_machine_warnings": len(new_issues),
        "question_answer_objects_unchanged": True,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
