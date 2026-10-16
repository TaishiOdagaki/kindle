"""Regression: replacements may retain old text while preserving adjacent runs."""
from docx import Document
from v9_apply_batch1 import replace_paragraph
doc=Document();p=doc.add_paragraph();p.add_run('1. ');p.add_run('内');p.add_run('容');p.add_run(' tail')
assert replace_paragraph(p,'内容','内容（Content）')
assert p.text=='1. 内容（Content） tail'
assert not replace_paragraph(p,'not present','other')
p=doc.add_paragraph('same same')
try:replace_paragraph(p,'same','new')
except AssertionError:pass
else:raise AssertionError('ambiguous duplicate must fail')
print('replacement: retained-substring, split runs, absent target, duplicate rejection passed')
