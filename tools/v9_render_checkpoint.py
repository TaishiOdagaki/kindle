"""Render an isolated Japanese-font QA copy; never alter the working DOCX.

Use the documents skill marker before authoring, then run with the primary
runtime Python. Supply a full CJK font with all manuscript CJK characters and a
writable copy of the bundled LibreOffice with its Noto Serif fallback. Webfont
subsets may be insufficient. Inspect the resulting PNGs manually.
"""
import argparse,importlib.util,sys,unicodedata
from pathlib import Path
from fontTools.ttLib import TTFont
from fontTools.varLib.instancer import instantiateVariableFont
from docx import Document
from docx.oxml.ns import qn

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--input',type=Path,required=True)
 ap.add_argument('--output',type=Path,required=True);ap.add_argument('--font',type=Path,required=True)
 ap.add_argument('--soffice',type=Path,required=True)
 ap.add_argument('--renderer',type=Path,default=Path('/root/.codex/skills/builtins/documents/render_docx.py'))
 args=ap.parse_args();args.output.mkdir(parents=True,exist_ok=True)
 assert args.input.is_file() and args.font.is_file() and args.soffice.is_file()
 # Never allow stale page images to be mistaken for this render.
 assert not list(args.output.glob('page-*.png')),'Use a fresh QA output directory'
 source_font=TTFont(args.font)
 doc=Document(args.input)
 # Japanese variable webfont subsets can omit simplified Chinese example
 # characters. Reject incomplete CJK coverage before any rendering is used.
 text=''.join(p.text for p in doc.paragraphs)
 for table in doc.tables:
  for row in table.rows:
   for cell in row.cells:text+=cell.text
 cjk={ord(c) for c in text if 0x3040<=ord(c)<=0x30ff or 0x3400<=ord(c)<=0x9fff or 0x20000<=ord(c)<=0x2ffff}
 missing=cjk-set(source_font.getBestCmap())
 assert not missing,'Incomplete CJK font coverage: '+''.join(chr(n) for n in sorted(missing))
 print(f'CJK font coverage: {len(cjk)} unique characters; missing 0')
 # IPA and other non-CJK characters may use LibreOffice's bundled fallback.
 # Verify availability as well as coverage; CJK coverage cannot prove this.
 fallback_path=args.soffice.resolve().parent.parent/'share/fonts/truetype/NotoSerif-Regular.ttf'
 assert fallback_path.is_file(),'Bundled Noto Serif fallback is required'
 fallback=TTFont(fallback_path)
 visible={ord(c) for c in text if not unicodedata.category(c).startswith(('C','Z'))}
 uncovered=visible-set(source_font.getBestCmap())-set(fallback.getBestCmap())
 assert not uncovered,'Incomplete fallback coverage: '+''.join(chr(n) for n in sorted(uncovered))
 print('Non-CJK fallback characters: '+''.join(chr(n) for n in sorted(visible-set(source_font.getBestCmap()))))
 font=instantiateVariableFont(source_font,{'wght':400},inplace=True) if 'fvar' in source_font else source_font
 for n in font['name'].names:
  if n.nameID in (1,4,6,16):
   n.string=('V9CJKQAStatic' if n.nameID==6 else 'V9 CJK QA Static').encode(n.getEncoding())
 fontdir=args.soffice.resolve().parent.parent/'share/fonts/truetype'
 fontdir.mkdir(parents=True,exist_ok=True);font.save(fontdir/'V9CJKQAStatic.ttf')
 for s in doc.styles:
  if s.type in (1,2):
   s.font.name='V9 CJK QA Static'
   s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'V9 CJK QA Static')
 for p in doc.paragraphs:
  for r in p.runs:
   r.font.name='V9 CJK QA Static'
   r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'V9 CJK QA Static')
 qa=args.output/'qa.docx';assert qa.resolve()!=args.input.resolve();doc.save(qa)
 spec=importlib.util.spec_from_file_location('documents_renderer',args.renderer)
 renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
 renderer._resolve_soffice=lambda:str(args.soffice.resolve())
 sys.argv=['render_docx.py',str(qa),'--output_dir',str(args.output),'--dpi','100','--emit_pdf']
 renderer.main()

if __name__=='__main__':main()
