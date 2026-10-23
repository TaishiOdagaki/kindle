"""Incremental 2027 precision fixes and explicit default typography."""
import copy,json,zipfile,hashlib
from pathlib import Path
from lxml import etree as E
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from apply_2027_round38 import docx_records,equivalent,rewrite_word,rewrite_xml
from audit_2027_epub import read_epub
R=Path(__file__).resolve().parents[1];O=R/'publishing/2027/round40';O.mkdir(exist_ok=True)
S=R/'publishing/2027/round38';wp=S/'2027_round38_review.docx';ep=wp.with_suffix('.epub')
d,W,wn=docx_records(wp);z,roots,Q,en,qa=read_epub(ep);B={q['id']:q for q in Q};WM={q['id']:q for q in W}
assert all(equivalent(q,WM[q['id']]) for q in Q)
M={q['id']:q for q in json.loads((R/'audit/v9/academic_master_500_20261022.json').read_text())['questions']};changes=[]
def change(id,reason,**fields):
 a=copy.deepcopy(B[id]);a.update(fields);changes.append(dict(id=id,before=B[id],after=a,reason=reason))
for id,reason in [('Q_2_012','必要に応じた意味の説明を正答にも復元。'),('Q_3_021','自分より有能な仲間という比較基準を復元。原著抜粋86頁を再取得。'),('Q_5_061','様態そうだの一用法を定義全体としないよう「には…用法があり」を復元。基金公式説明再照合。'),('Q_M_055','敬語の指針印刷40頁の「事実や気持ち」を復元。客観的恩恵だけに限定しない。')]:
 opts=list(B[id]['options']);k=B[id]['answer']-1;opts[k]=M[id]['options'][k];change(id,reason,options=opts)
change('Q_5_027','日本語発音教材内の説明に範囲を限定する。',stem='東京外国語大学の日本語発音教材の説明に照らし、日本語のアクセントとイントネーションの区別について、最も適切なものを一つ選びなさい。')
x=list(B['Q_4_083']['explanations']);x[0]='授業中に、その授業の状況に働き掛けながら判断や行動を見直す点で「行為の中の省察」の例です。ただし原著62頁のaction-presentは、実践の単位によって長さが異なります。授業の終了時刻だけで両者を機械的に分ける趣旨ではありません。';x[3]='原著61頁は、経験した実践を後から振り返り、そのときの理解や判断を検討して将来の実践に備える場合を述べています。本問は終了した授業を振り返る架空の例です。記録を基に判断の理由を検討する点で、単なる記録や同じ手順の反復とは区別されます。'
change('Q_4_083','Schön1983原著61〜62頁を本文・画像再照合。事後の振り返りとaction-presentの幅を明示。',explanations=x)
changed=set()
for e in changes:
 id=e['id'];b=e['before'];a=e['after']
 for f in ['stem','options','explanations']:
  if b[f]==a[f]:continue
  vals=[a[f]] if f=='stem' else a[f];assert len(vals)==len(wn[id][f])==len(en[id][f])
  for k,(v,p,t) in enumerate(zip(vals,wn[id][f],en[id][f]),1):
   if f!='stem' and b[f][k-1]==a[f][k-1]:continue
   lab=str(k)+'（'+('正' if k==a['answer'] else '誤')+'）' if f=='explanations' else None
   if f=='options':v=str(k)+'. '+v
   rewrite_word(p,v,lab);rewrite_xml(t,v,lab)
   for n,r in roots.items():
    if t.getroottree().getroot() is r:changed.add(n)
old=d.paragraphs[5].text;new=old.replace('設問・解説の学術的な独立再監査は継続中です。','');assert old!=new;rewrite_word(d.paragraphs[5],new);hits=0
for n,r in roots.items():
 for p in r.findall('.//{http://www.w3.org/1999/xhtml}p'):
  if ''.join(p.itertext()).strip()==old:rewrite_xml(p,new);changed.add(n);hits+=1
assert hits==1
s=d.styles['Normal'] if 'Normal' in d.styles else d.styles.add_style('Normal',WD_STYLE_TYPE.PARAGRAPH)
s._element.set(qn('w:default'),'1');s.font.name='Hiragino Mincho ProN';s.font.size=Pt(10.5);s.paragraph_format.line_spacing=Pt(14.5);s.paragraph_format.space_after=Pt(3)
f=s.element.get_or_add_rPr().find(qn('w:rFonts'))
for a in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+a),'Hiragino Mincho ProN')
for s in d.styles:
 if s.type==WD_STYLE_TYPE.PARAGRAPH and s.name.startswith('Heading'):
  s.font.name='Hiragino Sans';s.paragraph_format.line_spacing=1.0;s.paragraph_format.keep_with_next=True;f=s.element.get_or_add_rPr().find(qn('w:rFonts'));f.set(qn('w:eastAsia'),'Hiragino Sans')
d.styles['Title'].paragraph_format.line_spacing=1.15
for p in d.paragraphs:
 if p.style is not None and p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
 if p.style is not None and p.style.name=='Heading 1':p.paragraph_format.page_break_before=True
outw=O/'2027_review_manuscript.docx';oute=outw.with_suffix('.epub');d.save(outw)
with zipfile.ZipFile(oute,'w') as oz:
 for i in z.infolist():oz.writestr(i,E.tostring(roots[i.filename],encoding='utf-8',xml_declaration=True) if i.filename in changed else z.read(i.filename))
nd,NW,_=docx_records(outw);_,_,N,_,nqa=read_epub(oute);NM={q['id']:q for q in N};NWM={q['id']:q for q in NW};expected=copy.deepcopy(B)
for e in changes:expected[e['id']]=e['after']
assert NM==expected and all(equivalent(NM[i],NWM[i]) for i in NM)
assert all(q['answer']==NM[q['id']]['answer'] for q in Q)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};oldx=E.fromstring(zipfile.ZipFile(wp).read('word/document.xml'));newx=E.fromstring(zipfile.ZipFile(outw).read('word/document.xml'))
for query in ['//w:bookmarkStart/@w:name','//w:hyperlink/@w:anchor']:assert oldx.xpath(query,namespaces=ns)==newx.xpath(query,namespaces=ns)
assert len(nd.paragraphs)==7329 and len(nd.tables)==10
rep=dict(changes=changes,frontmatter_change=dict(before=old,after=new),layout='Added missing default Normal10.5pt Hiragino Mincho ProN, fixed14.5pt line after3pt; title and headings automatic line height; headings Hiragino Sans;10 major starts preserved',questions=500,word_epub_match=500,keys_changed=0,epub_qa=nqa,docx_sha256=hashlib.sha256(outw.read_bytes()).hexdigest(),epub_sha256=hashlib.sha256(oute.read_bytes()).hexdigest(),release_ready=False)
(O/'revisions_and_regression.json').write_text(json.dumps(rep,ensure_ascii=False,indent=2)+'\n');print('6 question precision changes;500 full-field matches;0 key changes')
