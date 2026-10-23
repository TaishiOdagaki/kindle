"""Final formatting only: preserve all Word text and reference targets."""
import re,json
from pathlib import Path
from docx import Document
from docx.shared import Pt
R=Path(__file__).resolve().parents[1];O=R/'publishing/2027/round40';p=O/'2027_review_manuscript.docx';d=Document(p)
before=[x.text for x in d.paragraphs];cases=[];gloss=[];inside=False
for i,x in enumerate(d.paragraphs):
 if re.match(r'^ケース\d+\u3000',x.text) and x.style.name=='Normal':
  x.paragraph_format.keep_with_next=True;cases.append(i)
 if x.style.name=='Heading 1':inside=x.text=='頻出用語集'
 elif inside:
  x.paragraph_format.space_after=Pt(2);gloss.append(i)
assert len(cases)==15
assert before==[x.text for x in d.paragraphs]
d.save(p)
(O/'layout_final_repairs.json').write_text(json.dumps(dict(case_headings_kept_with_next=cases,glossary_paragraphs_spacing_2pt=gloss,text_changed=False),ensure_ascii=False,indent=2)+'\n')
print(len(cases),len(gloss))
