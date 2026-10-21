#!/usr/bin/env python3
"""Clarify the 49-item appendix map and run a full structural cross-audit."""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "audit/v9"
DOCX = OUT / "登録日本語教員試験決定版問題集_v9_監査作業中.docx"
EXPECTED_PREVIOUS_DOCX = "966e1c22c2be017b191c4ae50f23ec1fd4e6a6905bc8a086c77b7dc7bcff1329"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def append_report(path: Path, text: str) -> None:
    old = path.read_text()
    marker = "## 第212チェックポイント 2026-10-21"
    assert marker not in old
    path.write_text(old.rstrip() + "\n\n\n" + text.strip() + "\n")


def structural_audit(path: Path) -> dict:
    sys.path.insert(0, str(ROOT / "tools"))
    import quiz_audit as qa

    issues, questions, answers = qa.audit(str(path))
    doc = Document(path)

    chapter_counts = {"1": 41, "2": 36, "3": 41, "4": 127, "5": 105,
                      "O": 75, "M": 60, "MO": 15}
    observed_counts = Counter(k.split("_")[1] for k in questions)
    assert dict(observed_counts) == chapter_counts
    assert len(questions) == len(answers) == 500

    table_ranges = [
        (0, "1", 1, 41), (1, "2", 1, 36), (2, "3", 1, 41),
        (3, "4", 1, 127), (4, "5", 1, 105), (5, "O", 1, 75),
        (6, "M", 1, 30), (7, "M", 31, 60), (8, "MO", 1, 15),
    ]
    answer_table_mismatches = []
    answer_rows = 0
    for table_index, prefix, start, end in table_ranges:
        rows = doc.tables[table_index].rows
        assert len(rows) == end - start + 2
        for row, number in zip(rows[1:], range(start, end + 1)):
            key = f"A_{prefix}_{number:03d}"
            actual = row.cells[1].text.strip()
            expected = answers[key]["ans"]
            answer_rows += 1
            if actual != expected:
                answer_table_mismatches.append({"key": key, "table": actual, "answer": expected})
    assert answer_rows == 500 and not answer_table_mismatches

    with zipfile.ZipFile(path) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
    bookmarks = [node.get(f"{{{W}}}name") for node in xml.xpath("//w:bookmarkStart", namespaces={"w": W})]
    anchors = [node.get(f"{{{W}}}anchor") for node in xml.xpath("//w:hyperlink", namespaces={"w": W})]
    bookmark_counts = Counter(bookmarks)
    anchor_counts = Counter(anchors)
    assert all(v == 1 for v in bookmark_counts.values())
    assert not (set(anchors) - set(bookmarks))
    for key in questions:
        answer_key = "A_" + key[2:]
        assert bookmark_counts[key] == bookmark_counts[answer_key] == 1
        assert anchor_counts[key] == anchor_counts[answer_key] == 1
    assert anchor_counts["TOC"] == 1000

    toc_expected = {
        "第1章　社会・文化・地域（41問）", "第2章　言語と社会（36問）",
        "第3章　言語と心理（41問）", "第4章　言語と教育（127問）",
        "第5章　言語（105問）", "教育実践ケーススタディ問題（75問）",
        "基礎試験60問＋応用試験15問（計75問）",
    }
    toc_found = {p.text.strip() for p in doc.paragraphs[:40] if "問）" in p.text}
    assert toc_expected <= toc_found

    case_titles = [p.text.strip() for p in doc.paragraphs if re.match(r"^ケース(?:[1-9]|1[0-5])　", p.text.strip())]
    assert len(case_titles) == 30
    assert set(Counter(case_titles).values()) == {2}

    appendix = doc.tables[9]
    appendix_text = " ".join(cell.text for row in appendix.rows for cell in row.cells)
    item_numbers = [int(x) for x in re.findall(r"<([0-9]+)>", appendix_text)]
    expanded = []
    for row in appendix.rows[1:]:
        text = row.cells[1].text
        ranges = [(int(a), int(b)) for a, b in re.findall(r"<([0-9]+)>[^<]*?〜<([0-9]+)>", text)]
        singles = [int(x) for x in re.findall(r"<([0-9]+)>", text)]
        expanded.extend(singles)
        for a, b in ranges:
            expanded.extend(range(a + 1, b))
    assert sorted(set(expanded)) == [n for n in range(1, 51) if n != 28]
    assert "養成課程コアカリキュラムの必須の教育内容（49項目）" in appendix.rows[0].cells[1].text
    assert "〈28〉教育実習は実践研修コアカリキュラムに位置付けられる" in appendix.rows[4].cells[1].text

    return {
        "questions": len(questions), "answers": len(answers),
        "chapter_counts": chapter_counts,
        "answer_quick_reference_rows": answer_rows,
        "answer_quick_reference_mismatches": answer_table_mismatches,
        "bookmarks": len(bookmarks), "unique_bookmarks": len(bookmark_counts),
        "internal_hyperlinks": len(anchors), "missing_anchor_targets": [],
        "question_answer_link_pairs": 500,
        "toc_return_links": anchor_counts["TOC"],
        "case_titles_problem_and_explanation": len(case_titles),
        "appendix_curriculum_items": 49,
        "appendix_excluded_item": 28,
        "raw_machine_warnings": len(issues),
    }


def main() -> None:
    assert digest(DOCX) == EXPECTED_PREVIOUS_DOCX
    before = Document(DOCX)
    old_header = "コアカリキュラムの必須の教育内容"
    new_header = "養成課程コアカリキュラムの必須の教育内容（49項目）"
    old_row = "<20>日本語教師の資質・能力〜<27>授業計画、<29>中間言語分析〜<36>著作権（全16項目）"
    new_row = (old_row + "。〈28〉教育実習は実践研修コアカリキュラムに位置付けられるため、"
               "養成課程の49項目には含めない。")
    assert before.tables[9].rows[0].cells[1].text == old_header
    assert before.tables[9].rows[4].cells[1].text == old_row

    additions = [
        {
            "id": "APPENDIX_MAP", "field": "table_cell", "index": None,
            "before": old_header, "after": new_header,
            "reason": "令和6年4月1日決定コアカリキュラムは、平成31年報告の〈28〉教育実習を実践研修側に位置付け、養成課程の必須の教育内容を49項目と明記する。表の対象と総数を明示する。",
            "sources": ["S163"], "applied": True,
            "table_index": 9, "row_index": 0, "column_index": 1,
        },
        {
            "id": "APPENDIX_MAP", "field": "table_cell", "index": None,
            "before": old_row, "after": new_row,
            "reason": "現行表は〈1〉〜〈50〉から〈28〉だけを除く49項目を正しく列挙するが、除外理由がなく脱落に見える。公式本文3頁の養成課程と実践研修の区分を明記する。",
            "sources": ["S163"], "applied": True,
            "table_index": 9, "row_index": 4, "column_index": 1,
        },
    ]
    changes_path = OUT / "changes.json"
    changes = json.loads(changes_path.read_text())
    assert not any(c.get("id") == "APPENDIX_MAP" and c.get("field") == "table_cell" for c in changes)
    changes.extend(additions)
    changes_path.write_text(json.dumps(changes, ensure_ascii=False, indent=2) + "\n")
    (OUT / "changes_batch212.json").write_text(json.dumps(additions, ensure_ascii=False, indent=2) + "\n")

    subprocess.run([sys.executable, str(ROOT / "tools/v9_rebuild.py")], check=True, cwd=ROOT / "tools")
    after = Document(DOCX)
    assert after.tables[9].rows[0].cells[1].text == new_header
    assert after.tables[9].rows[4].cells[1].text == new_row
    audit = structural_audit(DOCX)
    postfix = json.loads((OUT / "postfix_machine.json").read_text())
    assert postfix["replacements"] == 4323 and postfix["changed_questions"] == 493

    evidence = {
        "batch": 212, "date": "2026-10-21",
        "base_remote_commit": "112b04617460b66208287e952a9a9a8ae81caf46",
        "scope": "前付け、目次、500問リンク、正答早見表、ケース見出し、出題範囲対応表の構造横断点検。",
        "source": {
            "id": "S163", "title": "登録日本語教員 実践研修・養成課程コアカリキュラム",
            "decision_date": "2024-04-01", "pages_read": 15,
            "relevant_text": "本文3頁は〈28〉教育実習を実践研修コアカリキュラムに位置付け、養成課程の必須内容を49項目とする。",
        },
        "finding": "対応表の項目列挙は49項目で正しいが、〈28〉を飛ばす理由が書かれておらず、欠落に見える。",
        "changes_added": 2, "cumulative_changes": postfix["replacements"],
        "modified_questions": postfix["changed_questions"],
        "previous_docx_sha256": EXPECTED_PREVIOUS_DOCX,
        "current_docx_sha256": digest(DOCX),
        "structural_audit": audit,
        "question_answer_content_unchanged": True,
        "v8_sha256": digest(ROOT / "登録日本語教員試験決定版問題集_v8_修正版.docx"),
        "changed_pages_visual_qa_completed": False,
        "full_visual_qa_completed": False, "release_ready": False,
        "remaining": ["変更頁と隣接頁の再描画確認", "問題本文を含む全体の改頁調整", "最終機械検査", "最終全頁画像QA"],
    }
    assert evidence["v8_sha256"] == "4aef79631da7096e5837dd31473978d40d4932ba3cf8473c52138662a843d498"
    (OUT / "evidence/batch212_cross_audit.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")

    report = """## 第212チェックポイント 2026-10-21

前付け・目次・500問の問題／解説リンク・正答早見表・ケース見出し・出題範囲対応表を横断点検した。500問と500解答、500組の往復リンク、1000本の目次戻りリンク、正答早見表500行、章別問題数、15ケースの問題編／解説編見出しは全て一致した。

出題範囲対応表は〈1〉〜〈50〉のうち〈28〉だけを除く49項目を列挙していた。令和6年4月1日決定の登録日本語教員実践研修・養成課程コアカリキュラム全15頁を再照合し、〈28〉教育実習は実践研修側に位置付けられ、養成課程の必須内容は49項目であることを確認。表の脱落ではないが、理由がなく誤脱落に見えるため、表題と第4章欄へ区分を明記した。500問本文・正答・解説は変更していない。

変更頁の再描画確認、問題本文を含む全体の改頁調整、最終機械検査、最終全頁画像QAは未完了。最終版未制作、出版不可。詳細はevidence/batch212_cross_audit.json。"""
    append_report(OUT / "README.md", report)
    append_report(OUT / "critical_findings.md", report)
    append_report(OUT / "interim_report.md", report)
    append_report(OUT / "sources.md", """## 第212チェックポイント 2026-10-21

S163（令和6年4月1日決定「登録日本語教員 実践研修・養成課程コアカリキュラム」）全15頁を再照合。本文3頁は、平成31年報告の必須内容〈28〉教育実習を実践研修コアカリキュラムの内容とし、養成課程コアカリキュラムの必須の教育内容を49項目と明記する。巻末対応表は49項目を正しく列挙していたため、項目追加はせず、〈28〉を除く理由と49項目であることを明記した。""")

    manifest_path = OUT / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["full_visual_qa_completed"] = False
    manifest["release_ready"] = False
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"docx_sha256": digest(DOCX), **audit, "replacements": postfix["replacements"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
