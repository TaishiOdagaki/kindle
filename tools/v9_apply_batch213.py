#!/usr/bin/env python3
"""Apply and verify the global pagination policy for checkpoint 213."""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "audit/v9"
DOCX = OUT / "登録日本語教員試験決定版問題集_v9_監査作業中.docx"
EXPECTED_PREVIOUS_DOCX = "7ab038e3854af8d8b7d8697598c8b9047b956341d1f3df25553f5a6b5518a432"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit_pagination(path: Path) -> dict:
    doc = Document(path)
    question_lengths = {}
    probe = None
    for paragraph in doc.paragraphs:
        for bookmark in paragraph._element.iter(qn("w:bookmarkStart")):
            name = bookmark.get(qn("w:name"), "")
            if name.startswith(("Q_", "A_")):
                probe = name
        if probe and paragraph.text and probe.startswith("Q_"):
            question_lengths[probe] = question_lengths.get(probe, 0) + 1
        if probe and paragraph.text and "目次に戻る" in paragraph.text:
            probe = None
    assert len(question_lengths) == 500
    assert sorted(v for v in question_lengths.values() if v != 7) == [9, 10]
    context = None
    counts = {
        "question_paragraphs": 0, "question_keep_together": 0,
        "question_keep_with_next": 0, "answer_paragraphs": 0,
        "answer_keep_together": 0, "answer_keep_with_next": 0,
    }
    q_seen, a_seen = set(), set()
    position = 0
    for paragraph in doc.paragraphs:
        for bookmark in paragraph._element.iter(qn("w:bookmarkStart")):
            name = bookmark.get(qn("w:name"), "")
            if name.startswith(("Q_", "A_")):
                context = name
                position = 0
        text = paragraph.text
        if not context or not text:
            continue
        position += 1
        fmt = paragraph.paragraph_format
        if context.startswith("Q_"):
            q_seen.add(context)
            counts["question_paragraphs"] += 1
            counts["question_keep_together"] += fmt.keep_together is True
            length = question_lengths[context]
            chain = ("解説を見る" not in text and
                     (length != 10 or position == 1 or position >= length - 4))
            if chain:
                counts["question_keep_with_next"] += fmt.keep_with_next is True
        else:
            a_seen.add(context)
            counts["answer_paragraphs"] += 1
            counts["answer_keep_together"] += fmt.keep_together is True
            if "問題に戻る" not in text:
                counts["answer_keep_with_next"] += fmt.keep_with_next is True
        if "目次に戻る" in text:
            context = None
    assert len(q_seen) == len(a_seen) == 500
    assert counts["question_paragraphs"] == counts["question_keep_together"] == 3505
    assert counts["question_keep_with_next"] == 3001
    assert counts["answer_paragraphs"] == counts["answer_keep_together"]
    assert counts["answer_keep_with_next"] == 3000
    return counts | {"questions": len(q_seen), "answers": len(a_seen)}


def append_report(path: Path, text: str) -> None:
    old = path.read_text()
    marker = "## 第213チェックポイント 2026-10-21"
    assert marker not in old
    path.write_text(old.rstrip() + "\n\n\n" + text.strip() + "\n")


def main() -> None:
    assert digest(DOCX) == EXPECTED_PREVIOUS_DOCX
    sys.path.insert(0, str(ROOT / "tools"))
    import quiz_audit as qa

    before_issues, before_questions, before_answers = qa.audit(str(DOCX))
    subprocess.run([sys.executable, str(ROOT / "tools/v9_rebuild.py")], check=True,
                   cwd=ROOT / "tools")
    assert zipfile.ZipFile(DOCX).testzip() is None
    after_issues, after_questions, after_answers = qa.audit(str(DOCX))
    assert before_questions == after_questions
    assert before_answers == after_answers
    assert before_issues == after_issues
    counts = audit_pagination(DOCX)

    report = """## 第213チェックポイント 2026-10-21

全500問に共通の改頁規則を再構成工程へ実装した。通常問題は見出し・設問・4選択肢・解説リンクを一単位として連結し、最長の複数段落事例1問だけは見出しと冒頭、4選択肢と解説リンクをそれぞれ保持して事例本文の自然な段落境界でのみ改頁を許した。解説は見出し・正解・4肢解説・戻りリンクを一単位として連結した。これにより、選択肢や肢別解説だけの次頁送りと戻りリンクだけの頁頭孤立を防ぐ。本文、正答、選択肢、解説および変更履歴は変更していない。

OOXML検査は問題500、解説500、問題段落3505（全段落 keepLines、3001段落 keepNext）、解説段落3500（全段落 keepLines、ナビゲーション以外3000段落 keepNext）を確認した。DOCX再生成後も問題・解説の抽出結果と機械警告集合は再生成前と完全一致した。レンダリング検査結果は evidence/batch213_pagination.json に追記する。全頁画像目視は未完了のため、最終版未制作、出版不可である。"""
    for name in ("README.md", "critical_findings.md", "interim_report.md"):
        append_report(OUT / name, report)

    evidence = {
        "checkpoint": 213,
        "date": "2026-10-21",
        "purpose": "global pagination integrity without content changes",
        "pre_docx_sha256": EXPECTED_PREVIOUS_DOCX,
        "post_docx_sha256": digest(DOCX),
        "content_identity": {
            "questions_exact": before_questions == after_questions,
            "answers_exact": before_answers == after_answers,
            "machine_warnings_exact": before_issues == after_issues,
        },
        "ooxml_pagination": counts,
        "render": {"completed": False},
        "full_visual_qa_completed": False,
        "release_ready": False,
    }
    (OUT / "evidence/batch213_pagination.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
