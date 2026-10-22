#!/usr/bin/env python3
"""Create a NON-FINAL 2026 academic-branch Word build; never KDP upload."""
import argparse, hashlib, html, json, re, sys, zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

TITLE="2026年登録日本語教員試験　直前対策５００問＋詳細な解説"
AUTHOR="全国【登録日本語教員試験】研究グループ"
OLD="登録日本語教員試験　決定版問題集　基礎から応用試験対策＋本番予想模擬試験: 厳選　全５００問　詳細な解説付き"
def pset(p,text):
    if p.text==text:return
    if p._element.xpath(".//w:hyperlink"):raise ValueError("Unexpected hyperlink in replacement")
    if not p.runs:p.add_run(text)
    else:
        p.runs[0].text=text
        for x in p.runs[1:]:x.text=""
    assert p.text==text

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--input",required=True,type=Path)
    parser.add_argument("--output",required=True,type=Path)
    parser.add_argument("--report",required=True,type=Path)
    parser.add_argument("--master",required=True,type=Path)
    a=parser.parse_args()
    src=Document(a.input);p=src.paragraphs
    assert len(p)>6000
    # Front matter: before Table of Contents and before the first Q bookmark.
    toc=next(i for i,x in enumerate(p) if x.text.strip()=="目次")
    intro=next(i for i,x in enumerate(p[:toc]) if x.text.strip()=="はじめに")
    use=next(i for i,x in enumerate(p[:toc]) if x.text.strip()=="本書の使い方")
    assert 0<intro<use<toc<200
    pset(p[0],"2026年登録日本語教員試験\n直前対策５００問＋詳細な解説")
    if intro>=2:
        pset(p[1],"基礎・応用〈読解〉の直前復習を中心とする四肢択一演習")
        pset(p[2],"全500問・各選択肢の解説付き")
    pre=[
        "本書は、2026年度登録日本語教員試験に向け、重要事項を四肢択一500問で学習するための独自の問題集です。",
        "各問には正答だけでなく、残る選択肢がなぜ適切でないかについても解説を付しています。正解番号の暗記だけに頼らず、問題文の条件と各選択肢の意味を確認してください。",
        "構成は分野別425問と直前総仕上げ演習75問（基礎60問・応用読解15問）です。実試験の問題数・時間配分を完全再現した模擬試験ではありません。音声による聴解演習も収録していません。",
        "本書は、旧題『"+OLD+"』に基づく改訂版として制作しています。制度や試験実施要項は受験時点の公式発表を参照してください。",
        "著者表記：全国【登録日本語教員試験】研究グループ。活動経歴や所属等の事実は、別途著者が確認した情報に基づきます。"
    ]
    for j in range(intro+1,use):
        pset(p[j],pre[j-intro-1] if j-intro-1<len(pre) else "")
    directions=[
        "まず問題編を解き、解説を確認して、4つの選択肢それぞれの理由を確かめます。",
        "誤答だけでなく、正解していても根拠を説明できない問題を重点的に復習してください。",
        "問題下の「解説を見る」から解説へ移動し、解説末尾のリンクから問題へ戻れます。",
        "末尾の直前総仕上げ演習75問は短縮演習であり、本試験全科目を再現するものではありません。"
    ]
    for j in range(use+1,toc):
        pset(p[j],directions[j-use-1] if j-use-1<len(directions) else "")
    # Rewrite structural headings only. Do not touch the 500 question bodies.
    for x in p[toc:]:
        t=x.text
        if "模擬試験" in t and ("第III部" in t or "模擬試験"==t or "模擬試験　解説"==t):
            pset(x,t.replace("模擬試験","直前総仕上げ演習"))
    meta=src.core_properties;meta.title=TITLE;meta.author=AUTHOR
    meta.subject="2026年度登録日本語教員試験の基礎・応用〈読解〉を中心とした500問演習"
    meta.keywords="登録日本語教員試験,2026年,直前対策,500問"
    meta.comments="学術監査稿に基づく出版前候補。最終出版QAは未承認。"
    a.output.parent.mkdir(parents=True,exist_ok=True);src.save(a.output)
    doc=Document(a.output);master=json.loads(a.master.read_text())["questions"]
    assert len(master)==500
    for x in master:
        assert len(x["options"])==len(x["explanations"])==4
        assert x["answer"]==x["academic_review"]["independent_answer"]
    bookmarks=[]
    for para in doc.paragraphs:
        for bm in para._element.iter(qn("w:bookmarkStart")):
            name=bm.get(qn("w:name"),"")
            if re.fullmatch(r"[QA]_(?:[1-5]|O|M|MO)_\d{3}",name):bookmarks.append(name)
    assert len(bookmarks)==1000 and len(set(bookmarks))==1000,(len(bookmarks),len(set(bookmarks)))
    assert {x["id"] for x in master}=={x for x in bookmarks if x.startswith("Q_")}
    assert {"A_"+x["id"][2:] for x in master}=={x for x in bookmarks if x.startswith("A_")}
    with zipfile.ZipFile(a.output) as z:assert z.testzip() is None
    report={"title":TITLE,"source_word":str(a.input),"created_word":str(a.output),
        "sha256":hashlib.sha256(a.output.read_bytes()).hexdigest(),
        "academic_master_questions":len(master),"bookmarks":len(bookmarks),
        "full_500_field_content_match_NOT_tested":True,
        "full_visual_qa_completed":False,"epubcheck_completed":False,
        "cover_created":False,"release_ready":False,
        "warning":"Do not upload to KDP. A content-field comparison, rendering/visual QA, EPubCheck and cover/metadata checks remain."}
    a.report.parent.mkdir(parents=True,exist_ok=True)
    a.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
