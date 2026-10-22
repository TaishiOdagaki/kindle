#!/usr/bin/env python3
"""Safely assemble a *review candidate*, never an auto-approved Kindle release."""
import argparse,collections,hashlib,json,re,zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ID_RE=re.compile(r'[QA]_(?:[1-5]|O|M|MO)_\d{3}')
OPT_RE=re.compile(r'^\s*([1-4])[.．]\s+')
RAT_RE=re.compile(r'^\s*([1-4])（(?:正|誤)）\s*')

def anchors(doc):
    marks=[]
    for i,p in enumerate(doc.paragraphs):
        for bm in p._element.iter(qn('w:bookmarkStart')):
            name=bm.get(qn('w:name'),'')
            if ID_RE.fullmatch(name): marks.append((name,i))
    return sorted(marks,key=lambda x:x[1])

def rewrite(p,new,log,tag):
    old=p.text
    if old==new:return
    if any(x.tag==qn('w:hyperlink') for x in p._element):
        raise ValueError("hyperlink contained in edited paragraph: "+tag)
    if not p.runs:p.add_run(new)
    else:
        p.runs[0].text=new
        for r in p.runs[1:]:r.text=''
    assert p.text==new,tag
    log.append({'field':tag,'before':old,'after':new})

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--layout',required=True,type=Path)
    ap.add_argument('--academic-master',required=True,type=Path)
    ap.add_argument('--supplemental',type=Path)
    ap.add_argument('--output',required=True,type=Path)
    ap.add_argument('--report',required=True,type=Path)
    args=ap.parse_args()
    data=json.loads(args.academic_master.read_text())
    qs=data.get('questions',data)
    Q={q['id']:q for q in qs} if isinstance(qs,list) else qs
    assert len(Q)==500 and all(len(x['options'])==4 and len(x['explanations'])==4 for x in Q.values())
    doc=Document(args.layout); marks=anchors(doc)
    expectedQ=set(Q);expectedA={'A_'+id[2:] for id in Q}
    assert len(marks)==1000
    assert {id for id,_ in marks if id.startswith('Q_')}==expectedQ
    assert {id for id,_ in marks if id.startswith('A_')}==expectedA
    edited=[]; unresolved=[]
    for z,(id,start) in enumerate(marks):
        end=marks[z+1][1] if z+1<len(marks) else len(doc.paragraphs)
        ps=doc.paragraphs[start:end]
        if id.startswith('Q_'):
            q=Q[id]
            p0=next((i for i,p in enumerate(ps) if OPT_RE.match(p.text) and OPT_RE.match(p.text).group(1)=='1'),None)
            if p0 is None or p0==1:raise ValueError('choice boundary: '+id)
            parts=q['stem'].split('\n')
            if len(parts)==p0-1:
                for k,part in enumerate(parts):
                    rewrite(ps[k+1],part,edited,id+':stem'+str(k+1))
            else:
                rewrite(ps[1],parts[0],edited,id+':stem-first')
                unresolved.append({'id':id,'field':'stem_continuation',
                  'reason':f'layout has {p0-1} stem paragraphs vs {len(parts)} academic paragraphs'})
            for k in range(4):
                p=ps[p0+k];m=OPT_RE.match(p.text)
                if not m or m.group(1)!=str(k+1):raise ValueError('option count: '+id)
                rewrite(p,f'{k+1}. {q["options"][k]}',edited,id+':option'+str(k+1))
        else:
            q=Q['Q_'+id[2:]]
            if not re.search(r'正解[：:]\s*[1-4]',ps[1].text):raise ValueError('answer line: '+id)
            rewrite(ps[1],'正解：'+str(q['answer']),edited,id+':answer')
            for k in range(4):
                p=ps[k+2];m=RAT_RE.match(p.text)
                if not m or m.group(1)!=str(k+1):raise ValueError('explanation: '+id)
                tag='正' if k+1==q['answer'] else '誤'
                rewrite(p,f'{k+1}（{tag}）　{q["explanations"][k]}',edited,id+':explanation'+str(k+1))
    if args.supplemental and args.supplemental.exists():
        supplements=json.loads(args.supplemental.read_text()).get('edits',[])
        for k,x in enumerate(supplements):
            unresolved.append({'id':x.get('id'),'field':x['field'],'change_index':k,
               'reason':'contextual case/glossary/table/continuation edit not auto-applied'})
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.report.parent.mkdir(parents=True,exist_ok=True)
    doc.save(args.output)
    check=Document(args.output);a=anchors(check)
    assert len(a)==1000 and len(check.paragraphs)==len(doc.paragraphs)
    with zipfile.ZipFile(args.output) as z:assert z.testzip() is None
    report={'status':'MERGE_DRAFT_NOT_FOR_PUBLICATION','questions':500,
        'modified_paragraphs':len(edited),'unresolved_contextual_count':len(unresolved),
        'unresolved_contextual_details':unresolved,
        'qa':{'bookmarks':len(a),'file_opened':True,'zip_crc_ok':True},
        'sha256':hashlib.sha256(args.output.read_bytes()).hexdigest(),
        'release_ready':False,'epubcheck_completed':False,'full_page_visual_qa_completed':False}
    args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='unresolved_contextual_details'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
