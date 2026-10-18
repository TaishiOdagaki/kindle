"""Record full CJK coverage recovery without changing academic approval."""
import hashlib,json,subprocess,unicodedata
from pathlib import Path
from fontTools.ttLib import TTFont
from docx import Document
from pypdf import PdfReader
p=Path('audit/v9');root=Path(__file__).resolve().parents[2];r=root/'qa168';doc=p/'登録日本語教員試験決定版問題集_v9_監査作業中.docx'
def sha(f):return hashlib.sha256(f.read_bytes()).hexdigest()
def previous(n):return subprocess.check_output(['git','show','HEAD:audit/v9/'+n])
assert doc.read_bytes()==previous(doc.name)
assert len(json.loads((p/'changes.json').read_text()))==3363
assert (p/'manifest.json').read_bytes()==previous('manifest.json')
D=Document(doc);text=''.join(x.text for x in D.paragraphs)
for table in D.tables:
 for row in table.rows:
  for cell in row.cells:text+=cell.text
cjk={ord(c) for c in text if 0x3040<=ord(c)<=0x30ff or 0x3400<=ord(c)<=0x9fff or 0x20000<=ord(c)<=0x2ffff}
oldfont=TTFont(root/'NotoSansJP.ttf');font=TTFont(root/'NotoSansCJKjp-VF.ttf');fallbackfile=root/'lo165/share/fonts/truetype/NotoSerif-Regular.ttf';fallback=TTFont(fallbackfile)
assert cjk-set(oldfont.getBestCmap())=={ord('书')} and not cjk-set(font.getBestCmap())
visible={ord(c) for c in text if not unicodedata.category(c).startswith(('C','Z'))}
needed=visible-set(font.getBestCmap());assert needed=={ord(c) for c in 'ŋɕɯɸ'} and not needed-set(fallback.getBestCmap())
reader=PdfReader(r/'qa.pdf');assert len(reader.pages)==354
embedded=subprocess.check_output(['pdffonts',str(r/'qa.pdf')]).decode();assert 'V9CJKQAStatic' in embedded and 'NotoSerif-Regular' in embedded
pages={74:'fresh-074.png',93:'ipa-093.png',250:'review-250.png',267:'chinese-267.png',270:'ipa-270.png'}
e=dict(batch=168,kind='QA font recovery, no manuscript or academic status changes',base_remote_commit=subprocess.check_output(['git','rev-parse','HEAD']).decode().strip(),old_font_sha256=sha(root/'NotoSansJP.ttf'),old_missing_cjk=['书'],full_font_url='https://raw.githubusercontent.com/notofonts/noto-cjk/main/Sans/Variable/TTF/NotoSansCJKjp-VF.ttf',full_font_sha256=sha(root/'NotoSansCJKjp-VF.ttf'),full_font_bytes=(root/'NotoSansCJKjp-VF.ttf').stat().st_size,cjk_characters=len(cjk),missing_cjk=[],fallback_characters=sorted(chr(c) for c in needed),fallback_font='NotoSerif-Regular.ttf',fallback_sha256=sha(fallbackfile),missing_visible_with_fallback=[],incomplete_subset_rejected_before_render=True,deficient_font_render_used_as_qa=False,docx_sha256=sha(doc),pdf_sha256=sha(r/'qa.pdf'),render_pages=354,visually_reviewed_pages=sorted(pages),page_sha256={str(n):sha(r/f) for n,f in pages.items()},latest_batch167_text_matches=json.loads((r/'matches167.json').read_text()),pdf_fonts=embedded,full_visual_qa_completed=False,release_ready=False,academic_status_counts={'確認済み':397,'修正必要':2,'要追加確認':101},remaining=103)
(p/'evidence/batch168_font_qa.json').write_text(json.dumps(e,ensure_ascii=False,indent=2)+'\n')
report='''

## 第168バッチ 描画フォント検査 2026-10-18

学術判定・本文は変更せず、確認済み397・残103問、500問、3363件の累積修正を維持。Google配布Noto Sans JPのサブセットでは、中国語例字「书」が未収録だった。過去の変更頁目視は当該頁に限った確認であり、このフォントでの描画を全頁最終QAに転用しない。

Noto公式リポジトリの完全版Noto Sans CJK JP可変フォントを取得し、原稿の日本語・漢字1448種類の欠落ゼロを検査。非CJK文字のうち同フォントにないŋ・ɕ・ɯ・ɸは、同梱Noto Serifの収録を確認し、PDFで実際の埋込みも確認した。全可視文字は両フォントの和集合で収録される。描画前にCJKとフォールバック収録を検査する停止条件を追加し、旧サブセットが「书」で描画前に拒否される負例も検査済み。

完全版CJKフォントで全原稿を354頁に新規描画。フォント・計量変更で過去408頁から改頁が変わったため、旧頁番号のQA記録をこの描画の全頁QAへ流用しない。74・250頁で最新Q_4_089の9文字列と設問/4肢・全4解説を再確認。93・270頁でIPA、267頁で「书」の描画を直接確認し、欠字なし。74頁画像に見えた段落連結は同一PDFを別名で再描画し、原寸画像で各肢が分離していることを再確認。本文DOCX、v8、manifest、decisions、changesの学術判定は変更していない。

全354頁の最終目視は未実施、最終版・公開可能ではない。原フォントと代替フォントのハッシュ、PDF・確認画像のハッシュ、埋込みフォント一覧はevidence/batch168_font_qa.json。次回以降の描画には全CJK対応フォントを用い、ページマップは必ず当該描画から再生成する。
'''
for n in ['visual_qa.md','interim_report.md']:
 f=p/n;assert '## 第168バッチ' not in f.read_text();f.write_text(f.read_text()+report)
f=p/'README.md';f.write_text(f.read_text()+'''\n\n## 描画フォントの再開条件\n\n第168バッチでNoto Sans JPのWebfontサブセットが中国語例字「书」を欠くことを検出。以後は完全版Noto Sans CJK JP等を使用する。取得先：https://raw.githubusercontent.com/notofonts/noto-cjk/main/Sans/Variable/TTF/NotoSansCJKjp-VF.ttf 。原バイトSHA256はevidence/batch168_font_qa.json。tools/v9_render_checkpoint.pyはCJK収録と同梱Noto Serifの非CJKフォールバックを検査し、不足なら描画前に停止する。完全版での最新描画は354頁、全頁最終QAは未了。旧408頁の座標は流用せず再探索する。\n''')
